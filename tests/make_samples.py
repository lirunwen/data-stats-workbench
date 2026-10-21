"""生成用于验收的示例文件（可重复运行，覆盖正常、异常、重复、双列金额与邮件附件）。"""
from __future__ import annotations

import sys
from datetime import date
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
SAMPLES.mkdir(parents=True, exist_ok=True)


def make_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "经营流水"
    headers = ["日期", "收支类型", "经营类别", "部门", "往来单位", "金额", "结算方式", "单据号", "备注"]
    ws.append(headers)
    rows = [
        ["2026-08-03", "收入", "产品销售收入", "销售部", "杭州青禾科技有限公司", 128000, "银行转账", "DEMO-TEST-YS-001", "合同首款"],
        ["2026-08-04", "支出", "原材料采购", "产品部", "盛丰原材料有限公司", 46800, "银行转账", "DEMO-TEST-FK-001", "8月采购"],
        ["2026-08-05", "收入", "项目服务收入", "产品部", "远山实业集团", 86000, "银行承兑", "DEMO-TEST-YS-002", "阶段验收"],
        ["2026-08-06", "支出", "市场推广费", "市场部", "拓新广告传媒", 18000, "网银支付", "DEMO-TEST-FK-002", "线上投放"],
        ["2026-08-07", "收入", "渠道分销收入", "销售部", "澄明软件", 39200, "银行转账", "DEMO-TEST-YS-003", "渠道结算"],
        ["2026-08-10", "支出", "差旅交通费", "销售部", "云和差旅服务", 3860, "企业卡", "DEMO-TEST-FK-003", "华东出差"],
        ["2026-08-12", "收入", "产品销售收入", "销售部", "新澜数字科技", 156800, "银行转账", "DEMO-TEST-YS-004", "设备销售回款"],
        ["2026-08-13", "支出", "房租水电", "行政部", "蓝湾物业服务", 28000, "银行转账", "DEMO-TEST-FK-004", "8月物业费"],
        ["2026-08-15", "收入", "其他业务收入", "财务部", "长风商贸", 7800, "现金", "DEMO-TEST-YS-005", "废料处置"],
        ["2026-08-16", "支出", "员工薪酬", "财务部", "", 96000, "银行代发", "DEMO-TEST-FK-005", "7月工资"],
        ["2026-08-18", "收入", "项目服务收入", "产品部", "明途教育科技", 64000, "银行转账", "DEMO-TEST-YS-006", "课程系统尾款"],
        ["2026-08-19", "支出", "办公行政费", "行政部", "嘉禾办公用品", 1260, "现金", "DEMO-TEST-FK-006", "办公耗材"],
        ["2026-08-20", "", "产品销售收入", "销售部", "恒越医疗器械", 93000, "银行转账", "DEMO-TEST-YS-007", "未写收支类型"],
        ["2026-08-21", "支出", "市场推广费", "市场部", "东岸包装材料", 5200, "银行转账", "DEMO-TEST-FK-007", "物料制作"],
        ["2026-08-22", "收入", "渠道分销收入", "销售部", "禾木食品", 27000, "银行转账", "DEMO-TEST-YS-008", "月度结算"],
        ["2026-08-23", "支出", "原材料采购", "产品部", "万景机电设备", 45200, "银行承兑", "DEMO-TEST-FK-008", "设备备件"],
        ["2026-08-24", "收入", "产品销售收入", "销售部", "星桥物流", 112000, "银行转账", "DEMO-TEST-YS-009", "物流设备回款"],
        ["2026-08-25", "支出", "差旅交通费", "市场部", "云和差旅服务", -800, "企业卡", "DEMO-TEST-FK-009", "负数金额应被拦截"],
        ["2026-08-26", "收入", "产品销售收入", "销售部", "杭州青禾科技有限公司", 128000, "银行转账", "DEMO-TEST-YS-001", "与第 2 行重复"],
        ["2026-08-27", "转账", "其他支出", "运营部", "同源信息技术", 6600, "银行转账", "DEMO-TEST-FK-010", "未知收支类型"],
        ["2026-08-28", "支出", "", "", "诚达物流", 4300, "银行转账", "DEMO-TEST-FK-011", "缺类别和部门"],
        ["2026年8月29日", "收入", "其他业务收入", "运营部", "一格设计事务所", "¥12,300.00", "银行转账", "DEMO-TEST-YS-010", "中文日期与货币符号"],
        ["", "支出", "办公行政费", "行政部", "华信印刷", 950, "现金", "DEMO-TEST-FK-012", "日期为空"],
        ["2026-08-31", "支出", "办公行政费", "行政部", "华信印刷", "abc", "现金", "DEMO-TEST-FK-013", "金额非法"],
    ]
    for item in rows:
        ws.append(item)
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174D3F")
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDEFGHI", [12, 10, 18, 10, 24, 12, 12, 20, 22]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("说明")
    ws2["A1"] = "这是数据统计工作台生成的验收示例文件。"
    ws2["A2"] = "包含正常数据、重复、空日期、非法金额、未知收支等情形。"
    wb.save(SAMPLES / "经营流水_示例_含异常.xlsx")


def make_split_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "收支双列"
    ws.append(["业务日期", "经营类别", "部门", "往来单位", "收入金额", "支出金额", "单据编号", "备注"])
    for row in [
        ["2026-09-01", "产品销售收入", "销售部", "青屿客户A", 158000, 0, "DEMO-SP-001", "收入"],
        ["2026-09-02", "原材料采购", "产品部", "供应商A", 0, 38000, "DEMO-SP-002", "支出"],
        ["2026-09-03", "项目服务收入", "产品部", "客户B", 76000, 0, "DEMO-SP-003", "收入"],
        ["2026-09-04", "员工薪酬", "财务部", "", 0, 88000, "DEMO-SP-004", "支出"],
        ["2026-09-05", "市场推广费", "市场部", "广告商A", 0, 12500, "DEMO-SP-005", "支出"],
        ["2026-09-06", "其他业务收入", "运营部", "客户C", 4500, 0, "DEMO-SP-006", "收入"],
    ]:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174D3F")
    for col, width in zip("ABCDEFGH", [12, 18, 10, 20, 14, 14, 14, 16]):
        ws.column_dimensions[col].width = width
    wb.save(SAMPLES / "收入支出双列_示例.xlsx")


def make_email():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "九月费用"
    ws.append(["日期", "收支类型", "费用类别", "部门", "供应商", "金额", "单据号", "备注"])
    for row in [
        ["2026-09-08", "支出", "市场推广费", "市场部", "同源信息技术", 9800, "MAIL-FK-001", "邮件附件导入"],
        ["2026-09-09", "支出", "差旅交通费", "产品部", "云和差旅服务", 2650, "MAIL-FK-002", "客户拜访"],
        ["2026-09-10", "支出", "办公行政费", "行政部", "嘉禾办公用品", 780, "MAIL-FK-003", "办公用品"],
        ["2026-09-11", "支出", "房租水电", "行政部", "蓝湾物业服务", 15000, "MAIL-FK-004", "九月物业"],
    ]:
        ws.append(row)
    attachment = SAMPLES / "_邮件附件_九月费用.xlsx"
    wb.save(attachment)

    message = EmailMessage()
    message["Subject"] = "九月费用明细，请汇总入账"
    message["From"] = "财务小周 <finance@example.local>"
    message["To"] = "经营工作台 <workbench@example.local>"
    message["Date"] = "Mon, 12 Oct 2026 09:20:00 +0800"
    message.set_content("你好，九月费用明细见附件。附件中的 Excel 可以直接导入经营工作台。\n\n谢谢。")
    message.add_attachment(
        (SAMPLES / "_邮件附件_九月费用.xlsx").read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="九月费用明细.xlsx",
    )
    (SAMPLES / "经营数据邮件.eml").write_bytes(message.as_bytes())
    attachment.unlink(missing_ok=True)


if __name__ == "__main__":
    make_workbook()
    make_split_workbook()
    make_email()
    print("示例文件已生成：", SAMPLES)
