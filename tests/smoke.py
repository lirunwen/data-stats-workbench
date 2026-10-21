#!/usr/bin/env python3
"""端到端冒烟测试：后端 API、Excel/CSV/邮件导入、判重、CRUD、导出、备份。

用法：
    python tests/smoke.py

测试会使用独立的临时数据目录，不会触碰 data/ 下的真实数据。
"""
from __future__ import annotations

import io
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
TEMP = Path(tempfile.mkdtemp(prefix="workbench-smoke-"))
os.environ["DATASTATS_DATA_DIR"] = str(TEMP / "data")
os.environ.pop("DATASTATS_INSTALL_ROOT", None)
os.environ.pop("DATASTATS_VERSION_DIR", None)
sys.path.insert(0, str(ROOT))

from backend import config, db  # noqa: E402
from backend.errors import AppError  # noqa: E402
from app import create_app  # noqa: E402

PASSED = 0
FAILED = 0


def check(condition, message):
    global PASSED, FAILED
    if condition:
        PASSED += 1
        print(f"  ✓ {message}")
    else:
        FAILED += 1
        print(f"  ✗ {message}")
        raise AssertionError(message)


def api_payload(response):
    try:
        return response.get_json()
    except Exception:
        return None


def upload(client, filenames: list[Path] | list[tuple[str, bytes]]):
    data = []
    for item in filenames:
        if isinstance(item, Path):
            data.append((io.BytesIO(item.read_bytes()), item.name))
        else:
            data.append((io.BytesIO(item[1]), item[0]))
    return client.post(
        "/api/imports/analyze",
        data={"files": data},
        content_type="multipart/form-data",
    )


def selection(source):
    return [{"source_id": source["id"], "mapping": source["suggested_mapping"]}]


def run():
    config.configure(TEMP / "data")
    db.init_db()
    app = create_app()
    app.config.update(TESTING=True)
    client = app.test_client()

    print("\n[1] 健康检查与空状态")
    response = client.get("/api/health")
    check(response.status_code == 200 and response.get_json()["ok"], "健康检查")
    bootstrap = api_payload(client.get("/api/bootstrap"))
    check(bootstrap["app"]["options"]["total"] == 0, "空数据库")
    check("version" in bootstrap["app"], "Bootstrap 包含版本信息")

    print("\n[2] 安全边界")
    response = client.post("/api/demo", headers={"Origin": "http://evil.example"})
    check(response.status_code == 403, "拒绝跨站写入")
    response = client.get("/api/records", headers={"Host": "evil.example"})
    check(response.status_code == 403, "拒绝非本机 Host")

    print("\n[3] Excel 批量导入：识别、校验、异常、重复")
    response = upload(client, [SAMPLES / "经营流水_示例_含异常.xlsx"])
    payload = api_payload(response)
    check(response.status_code == 200 and payload["ok"], "上传并解析 Excel")
    check(payload["source_count"] == 1, "忽略说明工作表，只识别数据表")
    source = payload["sources"][0]
    token = payload["session_id"]
    check(source["row_count"] == 24, "读取 24 行数据")
    check(source["header_row"] == 1, "表头行识别正确")
    mapping = source["suggested_mapping"]
    check(mapping["fields"]["date"] == 0 and mapping["amount"]["column"] == 5, "自动映射字段")
    preview = client.post(
        f"/api/imports/{token}/preview",
        json={"sources": selection(source), "duplicate_policy": "skip"},
    )
    preview_payload = api_payload(preview)
    check(preview.status_code == 200 and preview_payload["ok"], "预览校验成功")
    summary = preview_payload["summary"]
    check(summary["invalid"] == 5, f"识别 5 行异常（实际 {summary['invalid']}）")
    check(summary["duplicate"] == 1, "识别批内重复")
    check(summary["warning"] == 1, "识别缺类别提示")
    check(summary["importable"] == 18, "可导入 18 条")
    commit = client.post(
        f"/api/imports/{token}/commit",
        json={"sources": selection(source), "duplicate_policy": "skip"},
    )
    batch_payload = api_payload(commit)
    check(commit.status_code == 200 and batch_payload["ok"], "事务提交成功")
    batch = batch_payload["batch"]
    check(batch["imported_rows"] == 18, "新增 18 条")
    check(batch["duplicate_rows"] == 1, "重复 1 条")
    check(batch["invalid_rows"] == 5, "异常 5 条")
    check(batch["status"] == "partial", "批次状态为部分导入")
    check(batch["sources"][0]["error_report"] is True, "生成异常明细报告")
    check(batch["sources"][0]["duplicate_rows"] == 1, "来源重复统计一致")

    print("\n[4] 重复导入与更新策略")
    response = upload(client, [SAMPLES / "经营流水_示例_含异常.xlsx"])
    payload = api_payload(response)
    source = payload["sources"][0]
    token2 = payload["session_id"]
    preview_skip = api_payload(client.post(
        f"/api/imports/{token2}/preview",
        json={"sources": selection(source), "duplicate_policy": "skip"},
    ))
    check(preview_skip["summary"]["importable"] == 0, "第二次导入跳过策略无可导入")
    check(preview_skip["summary"]["duplicate"] == 19, "第二次导入识别 19 条重复")
    rejected = client.post(
        f"/api/imports/{token2}/commit",
        json={"sources": selection(source), "duplicate_policy": "skip"},
    )
    check(rejected.status_code == 422, "全部重复/异常时拒绝提交")
    preview_update = api_payload(client.post(
        f"/api/imports/{token2}/preview",
        json={"sources": selection(source), "duplicate_policy": "update"},
    ))
    check(preview_update["summary"]["updateable"] == 18, "更新策略识别 18 条可更新")
    check(preview_update["summary"]["importable"] == 18, "更新策略可处理 18 条")
    update_commit = api_payload(client.post(
        f"/api/imports/{token2}/commit",
        json={"sources": selection(source), "duplicate_policy": "update"},
    ))
    update_batch = update_commit["batch"]
    check(update_batch["updated_rows"] == 18 and update_batch["imported_rows"] == 0, "重复导入只更新，不新增")
    check(update_batch["sources"][0]["duplicate_rows"] == 19, "更新策略保留完整重复统计")

    print("\n[5] 导入记录与异常报告")
    batches = api_payload(client.get("/api/batches?page_size=10"))
    check(batches["total"] >= 2, "导入批次可查询")
    check(all("session_token" not in item for item in batches["items"]), "批次列表不暴露会话令牌")
    detail = api_payload(client.get(f"/api/batches/{batch['id']}"))
    check(detail["batch"]["source_count"] == 1, "批次详情")
    report = client.get(f"/api/batches/{batch['id']}/sources/s1/errors.csv")
    check(report.status_code == 200 and report.data.startswith(b"\xef\xbb\xbf"), "CSV 异常报告带 BOM")
    archive = client.get(f"/api/batches/{batch['id']}/errors.zip")
    check(archive.status_code == 200 and archive.data[:2] == b"PK", "异常明细 ZIP")
    export_session = api_payload(upload(client, [SAMPLES / "经营流水_示例_含异常.xlsx"]))
    export_source = export_session["sources"][0]
    preview_export = client.post(
        f"/api/imports/{export_session['session_id']}/preview-export",
        json={"sources": selection(export_source), "duplicate_policy": "skip"},
    )
    check(preview_export.status_code == 200 and preview_export.data.startswith(b"\xef\xbb\xbf"), "校验明细导出")

    print("\n[6] 手工增删改查与筛选")
    manual = {
        "business_date": "2026-10-18", "direction": "收入", "amount": "12,345.67",
        "category": "验收测试收入", "department": "财务部", "counterparty": "测试客户",
        "payment_method": "银行转账", "doc_no": "SMOKE-MANUAL-001", "note": "CRUD 测试",
    }
    created = client.post("/api/records", json=manual)
    created_payload = api_payload(created)
    check(created.status_code == 201, "新增记录")
    record_id = created_payload["record"]["id"]
    check(created_payload["record"]["amount_cents"] == 1234567, "金额按分精确存储")
    duplicate = client.post("/api/records", json=manual)
    check(duplicate.status_code == 409, "手工重复记录被拦截")
    updated = client.put(f"/api/records/{record_id}", json={**manual, "amount": "22222.20", "note": "更新"})
    check(updated.status_code == 200 and api_payload(updated)["record"]["amount_cents"] == 2222220, "修改记录")
    filtered = api_payload(client.get("/api/records?q=SMOKE-MANUAL-001&direction=income"))
    check(filtered["total"] == 1 and filtered["items"][0]["note"] == "更新", "关键字与条件筛选")
    paged = api_payload(client.get("/api/records?page=1&page_size=10&sort=amount_desc"))
    check(paged["page"] == 1 and len(paged["items"]) <= 10, "分页")
    delete = client.delete(f"/api/records/{record_id}")
    check(delete.status_code == 200, "删除记录")
    check(client.get(f"/api/records/{record_id}").status_code == 404, "删除后不可读取")
    ids = []
    for index in range(3):
        result = client.post("/api/records", json={**manual, "doc_no": f"SMOKE-BULK-{index}", "amount": 100 + index})
        ids.append(api_payload(result)["record"]["id"])
    bulk = client.post("/api/records/bulk-delete", json={"ids": ids})
    check(bulk.status_code == 200 and api_payload(bulk)["deleted"] == 3, "批量删除")

    print("\n[7] 汇总看板")
    dashboard = api_payload(client.get("/api/dashboard?period=year"))
    check(dashboard["ok"] and dashboard["dashboard"]["has_data"], "看板有数据")
    check(len(dashboard["dashboard"]["trend"]) == 12, "趋势固定 12 个月")
    check(dashboard["dashboard"]["kpis"][0]["label"] == "经营收入", "KPI 标签")
    month = api_payload(client.get("/api/dashboard?period=month"))
    check(month["dashboard"]["period"]["start"].endswith("-01"), "月度期间起点")

    print("\n[8] Excel/CSV 导出与模板")
    export = client.get("/api/records/export?format=xlsx&direction=expense")
    check(export.status_code == 200 and export.data[:2] == b"PK", "导出 Excel")
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(export.data))
    check(workbook.sheetnames == ["经营流水", "汇总"], "导出工作簿含汇总页")
    check(workbook["经营流水"].max_row > 1, "导出有数据行")
    csv_export = client.get("/api/records/export?format=csv")
    check(csv_export.status_code == 200 and csv_export.data.startswith(b"\xef\xbb\xbf"), "导出 CSV 带 BOM")
    template = client.get("/api/import-template")
    check(template.status_code == 200 and template.data[:2] == b"PK", "下载导入模板")
    template_book = load_workbook(io.BytesIO(template.data))
    check(template_book.sheetnames == ["经营流水", "填写说明"], "模板工作表")

    print("\n[9] 邮件与多格式")
    response = upload(client, [SAMPLES / "经营数据邮件.msg"])
    msg_payload = api_payload(response)
    check(response.status_code == 200 and msg_payload["source_count"] == 2, "MSG 附件与正文表格")
    check(msg_payload["sources"][0]["presence"] == "attachment", "识别 MSG 附件")
    check(msg_payload["sources"][0]["suggested_mapping"]["amount"]["mode"] == "split", "双列金额自动识别")
    check(msg_payload["sources"][1]["presence"] == "body", "识别邮件正文表格")
    eml = api_payload(upload(client, [SAMPLES / "经营数据邮件.eml"]))
    check(eml["source_count"] == 1 and eml["files"][0]["attachments"][0]["status"] == "imported", "EML 附件")
    pst = api_payload(upload(client, [("旧邮箱.pst", b"not-a-real-pst")]))
    check(pst["source_count"] == 0 and "另存为 .msg" in pst["files"][0]["message"], "PST 给出明确导出指引")

    # 旧版 xls 与 GBK CSV，以及标题行在第二行的表头识别
    import xlwt
    xls_path = TEMP / "旧格式.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("数据")
    for col, value in enumerate(["日期", "收支", "类别", "金额", "单据号"]):
        sheet.write(0, col, value)
    for row, values in enumerate([
        ["2026-09-20", "收入", "测试收入", 500, "XLS-001"],
        ["2026-09-21", "支出", "测试支出", 200, "XLS-002"],
    ], start=1):
        for col, value in enumerate(values):
            sheet.write(row, col, value)
    book.save(str(xls_path))
    xls_payload = api_payload(upload(client, [xls_path]))
    check(xls_payload["source_count"] == 1 and xls_payload["sources"][0]["row_count"] == 2, "读取 .xls")

    csv_path = TEMP / "gBK台账.csv"
    csv_path.write_bytes("日期,收支类型,类别,部门,金额,单据号\n2026-09-22,收入,测试收入,销售部,880,CSV-001\n".encode("gbk"))
    csv_payload = api_payload(upload(client, [csv_path]))
    check(csv_payload["source_count"] == 1 and csv_payload["sources"][0]["row_count"] == 1, "读取 GBK CSV")

    from openpyxl import Workbook
    title_book = Workbook()
    title_sheet = title_book.active
    title_sheet["A1"] = "2026 年经营流水汇总"
    title_sheet.append(["业务日期", "收支类型", "经营类别", "金额", "单据编号"])
    title_sheet.append(["2026-09-23", "收入", "测试收入", 990, "TITLE-001"])
    title_path = TEMP / "标题行.xlsx"
    title_book.save(title_path)
    title_payload = api_payload(upload(client, [title_path]))
    check(title_payload["sources"][0]["header_row"] == 2, "自动跳过标题行识别表头")

    print("\n[10] 备份、设置与演示数据")
    backup = client.post("/api/maintenance/backup")
    check(backup.status_code == 200 and len(backup.data) > 1000, "手动备份")
    settings = client.put("/api/settings", json={"org_name": "冒烟测试单位", "default_duplicate_policy": "skip"})
    check(settings.status_code == 200 and api_payload(settings)["settings"]["org_name"] == "冒烟测试单位", "保存设置")
    demo = api_payload(client.post("/api/demo"))
    check(demo["created"] > 0, "载入演示数据")
    demo_dashboard = api_payload(client.get("/api/dashboard?period=year"))["dashboard"]
    check(demo_dashboard["has_data"] and demo_dashboard["demo_count"] > 0, "演示数据进入看板")
    cleared = api_payload(client.delete("/api/demo"))
    check(cleared["deleted"] == demo["created"], "清除演示数据")

    print("\n[11] 健康与清理")
    audit = api_payload(client.get("/api/audit"))
    check(audit["ok"] and len(audit["items"]) > 0, "操作审计可查询")
    cleanup = api_payload(client.post("/api/maintenance/cleanup"))
    check(cleanup["ok"], "清理临时文件")
    print(f"\n结果：{PASSED} 项通过，{FAILED} 项失败")
    if FAILED:
        raise SystemExit(1)


if __name__ == "__main__":
    try:
        run()
    except Exception as exc:
        import traceback
        traceback.print_exc()
        raise SystemExit(1)
    finally:
        shutil.rmtree(TEMP, ignore_errors=True)
