"""生成给文员使用的标准导入模板（纯本地 openpyxl）。"""
from __future__ import annotations

import io
from datetime import date

from . import config


def build_import_template() -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "经营流水"
    sheet.sheet_view.showGridLines = False

    headers = ["业务日期", "收支类型", "经营类别", "部门", "往来单位", "金额（元）", "结算方式", "单据编号", "摘要/备注"]
    examples = [
        [f"{date.today().year}-08-03", "收入", "产品销售收入", "销售部", "示例客户有限公司", 128000, "银行转账", "YS-2026-0001", "合同首款到账"],
        [f"{date.today().year}-08-05", "支出", "市场推广费", "市场部", "示例广告传媒", 16800, "网银支付", "FK-2026-0001", "季度推广投放"],
        [f"{date.today().year}-08-06", "收入", "项目服务收入", "产品部", "示例科技有限公司", 86000, "银行承兑", "YS-2026-0002", "项目阶段验收款"],
    ]
    sheet.append(headers)
    for row in examples:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="174D3F")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10, color="26352F")
    line = Side(style="thin", color="E0E5DD")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=line)
    for row in sheet.iter_rows(min_row=2, max_row=4):
        for cell in row:
            cell.font = body_font
            cell.border = Border(bottom=line)
            cell.alignment = Alignment(vertical="center")
        row[5].number_format = "#,##0.00"
    for row in range(2, 5001):
        sheet.cell(row=row, column=1).number_format = "yyyy-mm-dd"
        sheet.cell(row=row, column=6).number_format = "#,##0.00"
    widths = [13, 10, 20, 12, 26, 15, 13, 17, 32]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28
    for row in range(2, 5):
        sheet.row_dimensions[row].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:I5000"

    direction_validation = DataValidation(
        type="list", formula1='"收入,支出"', allow_blank=False,
        showErrorMessage=True, errorTitle="收支类型", error="请选择“收入”或“支出”",
    )
    sheet.add_data_validation(direction_validation)
    direction_validation.add("B2:B5000")

    # 单列说明，不会在导入时被识别成数据表。
    help_sheet = workbook.create_sheet("填写说明")
    help_sheet.sheet_view.showGridLines = False
    help_sheet.column_dimensions["A"].width = 116
    lines = [
        ("数据统计工作台 · 经营流水导入模板", 18, True, "174D3F"),
        ("", 10, False, "26352F"),
        ("填写方法", 12, True, "2F6B5B"),
        ("1. 在「经营流水」工作表中从第 2 行开始填写，每行一笔收入或支出。示例行可以直接删除。", 10, False, "4E5D55"),
        ("2. 业务日期建议使用 Excel 日期格式；也支持 2026-08-03、2026/8/3、2026年8月3日 等写法。", 10, False, "4E5D55"),
        ("3. 收支类型只填写“收入”或“支出”；也可以导入后在工作台中改为整表固定收支。", 10, False, "4E5D55"),
        ("4. 金额填写正数，不要带“元”以外的说明；支持 1,280.00、1280 等写法。", 10, False, "4E5D55"),
        ("5. 经营类别、部门越完整，仪表盘的分类与部门分析越准确。", 10, False, "4E5D55"),
        ("6. 有单据编号时，重复导入会按“年份 + 收支 + 单据编号 + 金额 + 类别”判重，避免误合并明细行。", 10, False, "4E5D55"),
        ("7. 同一文件可以有多个工作表；导入时每个工作表都可以单独选择、映射或排除。", 10, False, "4E5D55"),
        ("", 10, False, "26352F"),
        ("字段说明", 12, True, "2F6B5B"),
        ("业务日期：必填，交易或记账日期。", 10, False, "4E5D55"),
        ("收支类型：必填，收入或支出；也可在导入向导中选择整表固定。", 10, False, "4E5D55"),
        ("经营类别：建议填写，留空会归入“未分类”。", 10, False, "4E5D55"),
        ("部门：建议填写，用于部门经营分析。", 10, False, "4E5D55"),
        ("往来单位、结算方式、单据编号、摘要/备注：选填，用于检索与追溯。", 10, False, "4E5D55"),
        ("", 10, False, "26352F"),
        ("金额也可以拆成两列：收入金额、支出金额。在导入向导的“金额列”中选择“收入金额 / 支出金额双列”即可。", 10, False, "7A6339"),
        ("本文件只在本机读取，不会上传到任何外部服务。", 10, False, "7A6339"),
    ]
    for index, (text, size, bold, color) in enumerate(lines, start=1):
        cell = help_sheet.cell(row=index, column=1, value=text)
        cell.font = Font(name="微软雅黑", size=size, bold=bold, color=color)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if index == 1:
            help_sheet.row_dimensions[index].height = 34
    workbook.properties.creator = config.APP_NAME
    workbook.properties.title = "经营流水导入模板"
    workbook.properties.subject = "本地离线导入模板"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), "数据统计工作台_经营流水导入模板.xlsx"
