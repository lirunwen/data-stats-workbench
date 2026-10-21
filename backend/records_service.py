"""经营流水查询、增删改查与 Excel 导出。"""
from __future__ import annotations

import io
from datetime import date
from typing import Any

from . import config, db
from .errors import AppError
from .fields import FIELD_DEFS, FIELD_LABELS
from .utils import (
    FieldError,
    build_fingerprint,
    clean_text,
    csv_safe,
    direction_label,
    new_id,
    normalize_direction,
    now_iso,
    parse_date_input,
    parse_money_cents,
)

SORT_OPTIONS = {
    "date_desc": "business_date DESC, created_at DESC",
    "date_asc": "business_date ASC, created_at ASC",
    "amount_desc": "amount_cents DESC, business_date DESC",
    "amount_asc": "amount_cents ASC, business_date DESC",
    "updated_desc": "updated_at DESC",
}

LIKE_FIELDS = (
    "business_date", "category", "department", "counterparty", "payment_method",
    "doc_no", "note", "source_file",
)


def _escape_like(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def parse_filters(args: Any, *, for_export: bool = False) -> dict:
    params = {
        "start": None,
        "end": None,
        "direction": None,
        "category": None,
        "department": None,
        "q": None,
        "source_type": None,
        "batch_id": None,
        "sort": "date_desc",
        "page": 1,
        "page_size": 20 if not for_export else 100000,
    }
    start = clean_text(args.get("start"))
    end = clean_text(args.get("end"))
    if start:
        try:
            params["start"] = parse_date_input(start, "开始日期")
        except FieldError as exc:
            raise AppError(str(exc), 400, "invalid_filter") from exc
    if end:
        try:
            params["end"] = parse_date_input(end, "结束日期")
        except FieldError as exc:
            raise AppError(str(exc), 400, "invalid_filter") from exc
    if params["start"] and params["end"] and params["start"] > params["end"]:
        raise AppError("开始日期不能晚于结束日期", 400, "invalid_filter")

    direction = clean_text(args.get("direction"))
    if direction in {"income", "expense"}:
        params["direction"] = direction
    category = clean_text(args.get("category"))
    params["category"] = category[:80] if category else None
    department = clean_text(args.get("department"))
    params["department"] = department[:80] if department else None
    q = clean_text(args.get("q"))
    params["q"] = q[:100] if q else None
    source_type = clean_text(args.get("source"))
    if source_type in {"manual", "import"}:
        params["source_type"] = source_type
    batch_id = clean_text(args.get("batch_id"))
    params["batch_id"] = batch_id[:80] if batch_id else None

    sort = clean_text(args.get("sort"))
    if sort in SORT_OPTIONS:
        params["sort"] = sort
    if not for_export:
        try:
            params["page"] = max(1, int(args.get("page") or 1))
        except (TypeError, ValueError):
            params["page"] = 1
        try:
            page_size = int(args.get("page_size") or 20)
        except (TypeError, ValueError):
            page_size = 20
        params["page_size"] = page_size if page_size in {10, 20, 50, 100} else 20
    return params


def build_where(filters: dict) -> tuple[str, list[Any]]:
    clauses = []
    values: list[Any] = []
    if filters.get("start"):
        clauses.append("business_date >= ?")
        values.append(filters["start"])
    if filters.get("end"):
        clauses.append("business_date <= ?")
        values.append(filters["end"])
    if filters.get("direction"):
        clauses.append("direction = ?")
        values.append(filters["direction"])
    if filters.get("category"):
        clauses.append("category = ?")
        values.append(filters["category"])
    if filters.get("department"):
        clauses.append("department = ?")
        values.append(filters["department"])
    if filters.get("source_type"):
        clauses.append("source_type = ?")
        values.append(filters["source_type"])
    if filters.get("batch_id"):
        clauses.append("batch_id = ?")
        values.append(filters["batch_id"])
    if filters.get("q"):
        like = f"%{_escape_like(filters['q'])}%"
        clauses.append("(" + " OR ".join(f"{field} LIKE ? ESCAPE '\\'" for field in LIKE_FIELDS) + ")")
        values.extend([like] * len(LIKE_FIELDS))
    return (" WHERE " + " AND ".join(clauses)) if clauses else "", values


def record_to_dict(row: Any) -> dict:
    item = dict(row)
    item["amount"] = round(item["amount_cents"] / 100, 2)
    item["direction_label"] = direction_label(item["direction"])
    item["source_label"] = "批量导入" if item["source_type"] == "import" else "手工录入"
    item["source_location"] = ""
    if item["source_file"]:
        location = item["source_file"]
        if item.get("source_sheet"):
            location += f" · {item['source_sheet']}"
        if item.get("source_row"):
            location += f" · 第 {item['source_row']} 行"
        item["source_location"] = location
    item.pop("fingerprint", None)
    return item


def list_records(filters: dict) -> dict:
    conn = db.get_db()
    where, values = build_where(filters)
    total = conn.execute(f"SELECT COUNT(*) FROM records{where}", values).fetchone()[0]
    totals_rows = conn.execute(
        f"SELECT direction, SUM(amount_cents) AS amount, COUNT(*) AS count FROM records{where} GROUP BY direction",
        values,
    ).fetchall()
    income = next((int(r["amount"] or 0) for r in totals_rows if r["direction"] == "income"), 0)
    expense = next((int(r["amount"] or 0) for r in totals_rows if r["direction"] == "expense"), 0)
    page_size = filters["page_size"]
    pages = max(1, (total + page_size - 1) // page_size)
    page = min(filters["page"], pages)
    offset = (page - 1) * page_size
    order = SORT_OPTIONS.get(filters.get("sort"), SORT_OPTIONS["date_desc"])
    rows = conn.execute(
        f"SELECT * FROM records{where} ORDER BY {order} LIMIT ? OFFSET ?",
        [*values, page_size, offset],
    ).fetchall()
    return {
        "items": [record_to_dict(row) for row in rows],
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": pages,
        "summary": {
            "income_cents": income,
            "expense_cents": expense,
            "net_cents": income - expense,
            "count": total,
        },
        "filters": filters,
    }


def get_filter_options() -> dict:
    conn = db.get_db()
    categories = [
        {"name": row["category"], "count": row["count"]}
        for row in conn.execute(
            "SELECT category, COUNT(*) AS count FROM records GROUP BY category ORDER BY count DESC, category ASC"
        )
    ]
    departments = [
        {"name": row["department"], "count": row["count"]}
        for row in conn.execute(
            "SELECT department, COUNT(*) AS count FROM records GROUP BY department ORDER BY count DESC, department ASC"
        )
    ]
    years = [
        row["year"]
        for row in conn.execute(
            "SELECT DISTINCT substr(business_date, 1, 4) AS year FROM records ORDER BY year DESC"
        )
    ]
    range_row = conn.execute(
        "SELECT MIN(business_date) AS min_date, MAX(business_date) AS max_date, COUNT(*) AS total, "
        "SUM(CASE WHEN is_demo = 1 THEN 1 ELSE 0 END) AS demo_count FROM records"
    ).fetchone()
    batch_count = conn.execute("SELECT COUNT(*) FROM batches").fetchone()[0]
    return {
        "categories": categories,
        "departments": departments,
        "years": years,
        "min_date": range_row["min_date"],
        "max_date": range_row["max_date"],
        "total": range_row["total"],
        "demo_count": range_row["demo_count"] or 0,
        "batch_count": batch_count,
    }


def get_record(record_id: str) -> dict:
    row = db.get_db().execute("SELECT * FROM records WHERE id = ?", (record_id,)).fetchone()
    if not row:
        raise AppError("记录不存在或已删除", 404, "record_not_found")
    return record_to_dict(row)


def _normalize_payload(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise AppError("提交的数据格式不正确", 400)
    errors: dict[str, str] = {}
    try:
        business_date = parse_date_input(payload.get("business_date"), "业务日期")
    except FieldError as exc:
        business_date = ""
        errors["business_date"] = str(exc)
    try:
        direction = normalize_direction(payload.get("direction"))
    except FieldError as exc:
        direction = ""
        errors["direction"] = str(exc)
    amount_cents = 0
    try:
        amount_cents, _ = parse_money_cents(payload.get("amount"))
    except FieldError as exc:
        errors["amount"] = str(exc)
    values = {}
    for item in FIELD_DEFS:
        key = item["key"]
        if key in {"date", "amount"}:
            continue
        raw = payload.get(key, "")
        text = clean_text(raw, multiline=(key == "note"))
        limit = config.FIELD_LIMITS.get(key, 120)
        if len(text) > limit:
            errors[key] = f"{item['label']}不能超过 {limit} 个字符"
        values[key] = text
    if not values.get("category"):
        values["category"] = "未分类"
    if not values.get("department"):
        values["department"] = "未指定"
    if errors:
        raise AppError("请检查表单中的必填项和格式", 422, "validation_error", {"fields": errors})
    fingerprint = build_fingerprint(
        date_text=business_date,
        direction=direction,
        amount_cents=amount_cents,
        category=values["category"],
        department=values["department"],
        counterparty=values["counterparty"],
        note=values["note"],
        doc_no=values["doc_no"],
    )
    return {
        "business_date": business_date,
        "direction": direction,
        "category": values["category"],
        "department": values["department"],
        "counterparty": values["counterparty"],
        "payment_method": values["payment_method"],
        "doc_no": values["doc_no"],
        "note": values["note"],
        "amount_cents": amount_cents,
        "fingerprint": fingerprint,
    }


def create_record(payload: dict) -> dict:
    data = _normalize_payload(payload)
    conn = db.get_db()
    duplicate = conn.execute(
        "SELECT id FROM records WHERE fingerprint = ?", (data["fingerprint"],)
    ).fetchone()
    if duplicate:
        raise AppError(
            "系统中已有内容相同的记录，未重复添加。可修改单据编号或备注后重试。",
            409, "duplicate_record",
            {"existing_id": duplicate["id"]},
        )
    record_id = new_id("r")
    timestamp = now_iso()
    try:
        with db.transaction() as tx:
            tx.execute(
                """INSERT INTO records(id, business_date, direction, category, department, counterparty,
                       payment_method, doc_no, note, amount_cents, fingerprint, source_type, batch_id,
                       source_file, source_sheet, source_row, is_demo, created_at, updated_at)
                   VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'manual', NULL, '', '', NULL, 0, ?, ?)""",
                (
                    record_id, data["business_date"], data["direction"], data["category"],
                    data["department"], data["counterparty"], data["payment_method"], data["doc_no"],
                    data["note"], data["amount_cents"], data["fingerprint"], timestamp, timestamp,
                ),
            )
            db.audit(
                "create", "record", record_id,
                f"新增{direction_label(data['direction'])}记录 {data['amount_cents'] / 100:.2f} 元",
                conn=tx,
            )
    except Exception as exc:
        if "UNIQUE constraint failed: records.fingerprint" in str(exc):
            raise AppError("系统中已有内容相同的记录，未重复添加", 409, "duplicate_record") from exc
        raise
    return get_record(record_id)


def update_record(record_id: str, payload: dict) -> dict:
    conn = db.get_db()
    current = conn.execute("SELECT * FROM records WHERE id = ?", (record_id,)).fetchone()
    if not current:
        raise AppError("记录不存在或已删除", 404, "record_not_found")
    data = _normalize_payload(payload)
    duplicate = conn.execute(
        "SELECT id FROM records WHERE fingerprint = ? AND id <> ?", (data["fingerprint"], record_id)
    ).fetchone()
    if duplicate:
        raise AppError(
            "修改后的内容与另一条记录相同，请检查单据编号或备注。",
            409, "duplicate_record", {"existing_id": duplicate["id"]},
        )
    timestamp = now_iso()
    try:
        with db.transaction() as tx:
            tx.execute(
                """UPDATE records SET business_date = ?, direction = ?, category = ?, department = ?,
                       counterparty = ?, payment_method = ?, doc_no = ?, note = ?, amount_cents = ?,
                       fingerprint = ?, updated_at = ? WHERE id = ?""",
                (
                    data["business_date"], data["direction"], data["category"], data["department"],
                    data["counterparty"], data["payment_method"], data["doc_no"], data["note"],
                    data["amount_cents"], data["fingerprint"], timestamp, record_id,
                ),
            )
            db.audit(
                "update", "record", record_id,
                f"修改为{direction_label(data['direction'])} {data['amount_cents'] / 100:.2f} 元",
                conn=tx,
            )
    except Exception as exc:
        if "UNIQUE constraint failed: records.fingerprint" in str(exc):
            raise AppError("修改后的内容与另一条记录相同", 409, "duplicate_record") from exc
        raise
    return get_record(record_id)


def delete_record(record_id: str) -> dict:
    conn = db.get_db()
    row = conn.execute("SELECT * FROM records WHERE id = ?", (record_id,)).fetchone()
    if not row:
        raise AppError("记录不存在或已删除", 404, "record_not_found")
    with db.transaction() as tx:
        tx.execute("DELETE FROM records WHERE id = ?", (record_id,))
        db.audit(
            "delete", "record", record_id,
            f"删除{row['business_date']} {direction_label(row['direction'])} {row['amount_cents'] / 100:.2f} 元",
            conn=tx,
        )
    return {"deleted": 1}


def bulk_delete(ids: list[str]) -> dict:
    if not isinstance(ids, list) or not ids:
        raise AppError("请选择要删除的记录", 400)
    ids = [clean_text(i)[:80] for i in ids if clean_text(i)]
    if len(ids) > 500:
        raise AppError("一次最多删除 500 条记录", 400)
    marks = ",".join("?" for _ in ids)
    with db.transaction() as tx:
        count = tx.execute(f"SELECT COUNT(*) FROM records WHERE id IN ({marks})", ids).fetchone()[0]
        tx.execute(f"DELETE FROM records WHERE id IN ({marks})", ids)
        db.audit("bulk_delete", "record", "", f"批量删除 {count} 条经营流水", conn=tx)
    return {"deleted": count}


def export_filter_label(filters: dict) -> str:
    parts = []
    if filters.get("start") or filters.get("end"):
        parts.append(f"{filters.get('start') or '最早'} 至 {filters.get('end') or '最新'}")
    if filters.get("direction"):
        parts.append(direction_label(filters["direction"]))
    if filters.get("category"):
        parts.append(f"类别：{filters['category']}")
    if filters.get("department"):
        parts.append(f"部门：{filters['department']}")
    if filters.get("q"):
        parts.append(f"搜索：{filters['q']}")
    return "；".join(parts) if parts else "全部数据"


def build_export_rows(filters: dict) -> tuple[list[dict], dict]:
    conn = db.get_db()
    where, values = build_where(filters)
    total = conn.execute(f"SELECT COUNT(*) FROM records{where}", values).fetchone()[0]
    if total > 100000:
        raise AppError("本次导出超过 10 万条，请缩小日期范围后再导出", 400, "export_too_large")
    rows = conn.execute(
        f"SELECT * FROM records{where} ORDER BY {SORT_OPTIONS.get(filters.get('sort'), SORT_OPTIONS['date_desc'])}",
        values,
    ).fetchall()
    records = [record_to_dict(row) for row in rows]
    summary = {
        "count": len(records),
        "income_cents": sum(r["amount_cents"] for r in records if r["direction"] == "income"),
        "expense_cents": sum(r["amount_cents"] for r in records if r["direction"] == "expense"),
    }
    summary["net_cents"] = summary["income_cents"] - summary["expense_cents"]
    return records, summary


def _safe_excel_text(value: Any) -> str:
    text = clean_text(value, multiline=True)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def export_xlsx(filters: dict) -> tuple[bytes, str]:
    """导出为 Excel（流式写入）。

    必须使用 write_only 模式：openpyxl 的常规模式会为每个单元格复制样式代理，
    单行成本随行数上升（实测 200 行 0.35 ms/行，1 万行 2.1 ms/行，5 万行约 10 ms/行），
    5 万行导出会退化到 8 分钟以上。write_only 直接流式写 XML，同样数据只需数秒。

    代价是不再写入 Excel「表格对象」（Table）。已用「自动筛选 + 交替底色」替代，
    用户的筛选、排序、打印体验不变。
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    records, summary = build_export_rows(filters)

    # 样式对象只构造一次，全表复用；原实现每格新建，5 万行要构造 65 万次。
    header_fill = PatternFill("solid", fgColor="173F35")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10, color="2C3330")
    band_fill = PatternFill("solid", fgColor="FAFBF8")
    border = Border(bottom=Side(style="thin", color="E4E8E2"))
    header_align = Alignment(horizontal="center", vertical="center")
    body_align = Alignment(vertical="center")
    body_align_wrap = Alignment(vertical="center", wrap_text=True)
    wrap_columns = {10, 12}

    headers = ["序号", "业务日期", "收支类型", "经营类别", "部门", "往来单位", "金额（元）",
               "结算方式", "单据编号", "摘要 / 备注", "来源", "来源位置", "录入时间"]

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("经营流水")
    sheet.sheet_view.showGridLines = False
    # write_only 会在第一次 append 时写出工作表头部（cols / sheetView / autoFilter），
    # 因此列宽、冻结窗格、筛选范围必须在 append 之前设置，否则会被静默丢弃。
    widths = [6, 12, 9, 18, 12, 22, 14, 12, 16, 30, 10, 36, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    if records:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:M{len(records) + 1}"

    header_cells = []
    for value in headers:
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border
        header_cells.append(cell)
    sheet.row_dimensions[1].height = 28
    sheet.append(header_cells)

    for idx, record in enumerate(records, start=1):
        values = [
            idx,
            record["business_date"],
            record["direction_label"],
            _safe_excel_text(record["category"]),
            _safe_excel_text(record["department"]),
            _safe_excel_text(record["counterparty"]),
            record["amount_cents"] / 100,
            _safe_excel_text(record["payment_method"]),
            _safe_excel_text(record["doc_no"]),
            _safe_excel_text(record["note"]),
            record["source_label"],
            _safe_excel_text(record["source_location"]),
            record["created_at"].replace("T", " ")[:19],
        ]
        cells = []
        for position, value in enumerate(values, start=1):
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = body_font
            cell.border = border
            cell.alignment = body_align_wrap if position in wrap_columns else body_align
            if idx % 2 == 0:
                cell.fill = band_fill
            cells.append(cell)
        cells[6].number_format = "#,##0.00"
        # 行高同样要在 append 之前登记，否则不会被写入 XML。
        sheet.row_dimensions[idx + 1].height = 22
        sheet.append(cells)

    # ---- 汇总页 ----
    summary_sheet = workbook.create_sheet("汇总")
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.column_dimensions["A"].width = 20
    summary_sheet.column_dimensions["B"].width = 22
    summary_sheet.column_dimensions["C"].width = 18
    summary_sheet.column_dimensions["D"].width = 18
    # write_only 工作表没有 merge_cells 方法，直接向 merged_cells 登记区域即可。
    summary_sheet.merged_cells.add("A1:C1")

    title = WriteOnlyCell(summary_sheet, value="经营数据导出汇总")
    title.font = Font(name="微软雅黑", size=16, bold=True, color="173F35")
    summary_sheet.append([title])
    summary_sheet.append([None])  # 空行，使元信息从第 3 行开始

    label_font = Font(name="微软雅黑", size=10, bold=True, color="66716B")
    value_font = Font(name="微软雅黑", size=10, color="1D2622")
    meta_rows = [
        ("导出时间", date.today().isoformat()),
        ("筛选条件", export_filter_label(filters)),
        ("记录笔数", summary["count"]),
        ("收入合计（元）", summary["income_cents"] / 100),
        ("支出合计（元）", summary["expense_cents"] / 100),
        ("经营净额（元）", summary["net_cents"] / 100),
    ]
    for label, value in meta_rows:
        label_cell = WriteOnlyCell(summary_sheet, value=label)
        label_cell.font = label_font
        value_cell = WriteOnlyCell(summary_sheet, value=value)
        value_cell.font = value_font
        if "元" in label:
            value_cell.number_format = "#,##0.00"
        summary_sheet.append([label_cell, value_cell])

    summary_sheet.append([None])  # 空行，分类汇总表头固定在第 10 行
    header_fill_summary = PatternFill("solid", fgColor="2F6B5B")
    summary_align = Alignment(horizontal="center")
    category_header = []
    for value in ("经营类别", "收入（元）", "支出（元）", "净额（元）"):
        cell = WriteOnlyCell(summary_sheet, value=value)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill_summary
        cell.alignment = summary_align
        category_header.append(cell)
    summary_sheet.append(category_header)

    categories: dict[str, list[int]] = {}
    for record in records:
        bucket = categories.setdefault(record["category"], [0, 0])
        bucket[0 if record["direction"] == "income" else 1] += record["amount_cents"]
    for name, amounts in sorted(categories.items(), key=lambda item: sum(item[1]), reverse=True):
        name_cell = WriteOnlyCell(summary_sheet, value=_safe_excel_text(name))
        income_cell = WriteOnlyCell(summary_sheet, value=amounts[0] / 100)
        expense_cell = WriteOnlyCell(summary_sheet, value=amounts[1] / 100)
        net_cell = WriteOnlyCell(summary_sheet, value=(amounts[0] - amounts[1]) / 100)
        for cell in (income_cell, expense_cell, net_cell):
            cell.number_format = "#,##0.00"
        summary_sheet.append([name_cell, income_cell, expense_cell, net_cell])

    workbook.properties.creator = config.APP_NAME
    workbook.properties.title = "经营流水导出"
    output = io.BytesIO()
    workbook.save(output)
    filename = f"经营流水_{date.today().strftime('%Y%m%d')}.xlsx"
    return output.getvalue(), filename


def export_csv(filters: dict) -> tuple[bytes, str]:
    import csv

    records, _summary = build_export_rows(filters)
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(["业务日期", "收支类型", "经营类别", "部门", "往来单位", "金额（元）",
                     "结算方式", "单据编号", "摘要/备注", "来源", "来源位置"])
    for record in records:
        writer.writerow([
            record["business_date"],
            record["direction_label"],
            csv_safe(record["category"]),
            csv_safe(record["department"]),
            csv_safe(record["counterparty"]),
            f"{record['amount_cents'] / 100:.2f}",
            csv_safe(record["payment_method"]),
            csv_safe(record["doc_no"]),
            csv_safe(record["note"]),
            record["source_label"],
            csv_safe(record["source_location"]),
        ])
    filename = f"经营流水_{date.today().strftime('%Y%m%d')}.csv"
    return ("\ufeff" + output.getvalue()).encode("utf-8"), filename
